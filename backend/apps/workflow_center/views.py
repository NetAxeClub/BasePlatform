import json
import operator
from datetime import datetime
from io import BytesIO
from urllib.parse import quote
from uuid import uuid4

import requests
from django.apps import apps
from django.db.models import CharField, Count, ForeignKey, GenericIPAddressField
from django.http import HttpResponse, JsonResponse
from django.utils import timezone
from django_filters.rest_framework import DjangoFilterBackend
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from rest_framework import filters
from rest_framework.decorators import action
from rest_framework.views import APIView

from apps.api.tools.custom_pagination import LargeResultsSetPagination
from apps.api.tools.custom_viewset_base import CustomViewBase
from apps.workflow_center.inspection import inspection_scope, inspection_type
from apps.workflow_center.models import (
    Method,
    State,
    Tasks,
    WorkflowExecution,
    WorkflowHostVar,
    WorkflowInventory,
)
from apps.workflow_center.serializers import (
    WorkflowExecutionSerializer,
    WorkflowHostVarSerializer,
    WorkflowInventorySerializer,
)
from apps.workflow_center.services import (
    WorkflowDiagnoseService,
    get_address_trace_columns,
    get_device_interfaces,
    get_firewall_list,
    query_address_traces,
)
from confload.confload import config
from driver import auto_driver_map


class WorkflowExecutionViewSet(CustomViewBase):
    queryset = WorkflowExecution.objects.all().order_by("-commit_time")
    serializer_class = WorkflowExecutionSerializer
    filter_backends = (DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter)
    filter_fields = "__all__"
    ordering_fields = ("-id",)

    @staticmethod
    def _parse_datetime(value):
        if not value:
            return None
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
            try:
                dt = datetime.strptime(value, fmt)
                if fmt == "%Y-%m-%d":
                    dt = datetime(dt.year, dt.month, dt.day, 0, 0, 0)
                return timezone.make_aware(dt) if timezone.is_naive(dt) else dt
            except ValueError:
                continue
        return None

    @action(detail=False, methods=["get"])
    def inspection_results(self, request):
        queryset = self.get_queryset().filter(task=Tasks.INSPECTION)
        query_params = request.query_params

        db_filters = {
            "commit_user": query_params.get("commit_user"),
            "device": query_params.get("device"),
            "origin": query_params.get("origin"),
            "state": query_params.get("state"),
            "task_id": query_params.get("task_id"),
        }
        for field_name, value in db_filters.items():
            if value not in (None, ""):
                queryset = queryset.filter(**{field_name: value})

        if query_params.get("code") not in (None, ""):
            queryset = queryset.filter(code=query_params.get("code"))

        start_time = self._parse_datetime(query_params.get("start_time"))
        if start_time:
            queryset = queryset.filter(commit_time__gte=start_time)
        end_time = self._parse_datetime(query_params.get("end_time"))
        if end_time:
            queryset = queryset.filter(commit_time__lte=end_time)

        limit = int(query_params.get("limit", 100))
        flows = list(queryset.order_by("-commit_time")[:limit])

        scope_filter = query_params.get("inspection_scope", "").lower()
        type_filter = query_params.get("inspection_type", "")
        keyword = query_params.get("keyword", "").lower()

        filtered = []
        for flow in flows:
            flow_scope = inspection_scope(flow)
            flow_type = inspection_type(flow)
            if scope_filter and flow_scope != scope_filter:
                continue
            if type_filter and flow_type != type_filter:
                continue
            if keyword:
                haystack = " ".join(
                    [
                        str(getattr(flow, "task_id", "")),
                        str(getattr(flow, "commit_user", "")),
                        str(getattr(flow, "device", "")),
                        str(getattr(flow, "task_result", "")),
                        str(getattr(flow, "kwargs", "")),
                    ]
                ).lower()
                if keyword not in haystack:
                    continue
            filtered.append(flow)

        serializer = self.get_serializer(filtered, many=True)
        return JsonResponse({"code": 200, "msg": "success", "data": serializer.data, "count": len(serializer.data)})

    @action(detail=True, methods=["get"])
    def inspection_detail(self, request, *args, **kwargs):
        instance = self.get_object()
        if instance.task != Tasks.INSPECTION:
            return JsonResponse({"code": 400, "msg": "当前记录不是巡检结果", "data": None})
        serializer = self.get_serializer(instance)
        return JsonResponse({"code": 200, "msg": "success", "data": serializer.data})

    @action(detail=False, methods=["post"])
    def write_inspection_result(self, request):
        payload = request.data
        device = payload.get("device") or payload.get("device_ip")
        commit_user = (
            payload.get("commit_user")
            or getattr(getattr(request, "user", None), "username", "")
            or "system"
        )
        method = payload.get("method") or Method.RESTAPI
        state = payload.get("state") or State.FINISH
        code = payload.get("code")
        if code in (None, ""):
            code = 9008 if state == State.FINISH else 9006 if state == State.PUBLISHED else 9005 if state == State.FAILED else 9000

        kwargs_payload = payload.get("kwargs", {}) or {}
        inspection_meta = {
            "inspection_scope": payload.get("inspection_scope") or kwargs_payload.get("inspection_scope") or ("device" if device else "fleet"),
            "inspection_type": payload.get("inspection_type") or kwargs_payload.get("inspection_type") or "generic",
            "summary": payload.get("summary", {}),
        }
        kwargs_payload.update(inspection_meta)

        task_result = payload.get("task_result", {})
        ttp = payload.get("ttp", {})
        commands = payload.get("commands", [])
        back_off_commands = payload.get("back_off_commands", [])

        flow = WorkflowExecution.objects.create(
            task_id=payload.get("task_id") or f"inspection-{uuid4().hex[:12]}",
            origin=payload.get("origin", "NetClaw-CN"),
            task_result=json.dumps(task_result, ensure_ascii=False) if isinstance(task_result, (dict, list)) else str(task_result or ""),
            order_code=payload.get("order_code"),
            device=device,
            device_id=payload.get("device_id"),
            commit_user=commit_user,
            commit_time=timezone.now(),
            task=Tasks.INSPECTION,
            method=method,
            class_method=payload.get("class_method", "inspection_result"),
            remote_ip=str(request.META.get("REMOTE_ADDR", "")),
            kwargs=json.dumps(kwargs_payload, ensure_ascii=False),
            ttp=json.dumps(ttp, ensure_ascii=False) if isinstance(ttp, (dict, list)) else str(ttp or "{}"),
            commands=json.dumps(commands, ensure_ascii=False) if isinstance(commands, list) else str(commands or "[]"),
            back_off_commands=json.dumps(back_off_commands, ensure_ascii=False) if isinstance(back_off_commands, list) else str(back_off_commands or "[]"),
            state=state,
            code=code,
        )
        serializer = self.get_serializer(flow)
        return JsonResponse({"code": 201, "msg": "success", "data": serializer.data})


class WorkflowHostVarViewSet(CustomViewBase):
    queryset = WorkflowHostVar.objects.all().order_by("-id")
    queryset = WorkflowHostVarSerializer.setup_eager_loading(queryset)
    serializer_class = WorkflowHostVarSerializer
    filter_backends = (DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter)
    filter_fields = "__all__"
    pagination_class = LargeResultsSetPagination
    search_fields = ("name", "host")


class WorkflowInventoryViewSet(CustomViewBase):
    queryset = WorkflowInventory.objects.all().order_by("-id")
    queryset = WorkflowInventorySerializer.setup_eager_loading(queryset)
    serializer_class = WorkflowInventorySerializer
    filter_backends = (DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter)
    filter_fields = "__all__"
    pagination_class = LargeResultsSetPagination


class WorkflowAutomationChartView(APIView):
    def get(self, request):
        get_params = request.GET.dict()
        if "event_task_module" in get_params:
            data = list(WorkflowExecution.objects.values("task").annotate(sum_count=Count("task")))
            data = sorted(data, key=operator.itemgetter("sum_count"), reverse=True)
            return JsonResponse({"code": 200, "data": data}, safe=False)

        if "event_task_user" in get_params:
            data = list(WorkflowExecution.objects.values("commit_user").annotate(sum_count=Count("commit_user")))
            data = sorted(data, key=operator.itemgetter("sum_count"), reverse=True)
            return JsonResponse({"code": 200, "data": data}, safe=False)

        if "event_commit_time" in get_params:
            work_time_list = []
            event_commit_queryset = WorkflowExecution.objects.values().all()
            for item in event_commit_queryset:
                current_day = item["commit_time"].strftime("%Y-%m-%d %H:%M:%S")[0:11]
                if current_day + "08:30" < item["commit_time"].strftime("%Y-%m-%d %H:%M:%S") < current_day + "17:30":
                    work_time_list.append(item)
            return JsonResponse(
                {
                    "code": 200,
                    "data": {
                        "work_time_count": len(work_time_list),
                        "total_time_count": len(event_commit_queryset),
                        "not_work_time": len(event_commit_queryset) - len(work_time_list),
                    },
                },
                safe=False,
            )

        return JsonResponse({"code": 400, "data": []}, safe=False)


class WorkflowAddressLocationView(APIView):
    def get(self, request):
        get_param = request.GET.dict()
        mongo_data = {}

        if get_param.get("get_layer2interface_hostip"):
            hostip = get_param["get_layer2interface_hostip"]
            interface_str = get_param["interfaces"]
            interface_list = [iface.strip() for iface in interface_str.split(",") if iface.strip()]
            results = {}
            for interface in interface_list:
                query_string = f'ifHCInOctets{{instance="{hostip}", ifName="{interface}"}}'
                grafana_api = config.grafana_net_device_resource_api
                grafana_token = config.grafana_net_device_resource_api_token
                headers = {"Authorization": f"Bearer {grafana_token}"}
                response = requests.request("POST", grafana_api, headers=headers, params={"query": query_string})
                results[interface] = response.json()["data"]["result"][0]["metric"]["interface"]
            return JsonResponse({"code": 200, "results": results}, safe=False)

        if get_param.get("get_interface_by_hostip"):
            hostip = get_param["get_interface_by_hostip"]
            res = get_device_interfaces(hostip)
            return JsonResponse({"code": 200, "count": len(res), "message": "成功", "results": res}, safe=False)

        if get_param.get("get_table_columns") == "1":
            return JsonResponse({"code": 200, "results": get_address_trace_columns()}, safe=False)

        page_size = int(get_param.get("page_size") or get_param.get("limit") or 10)
        page_num = int(get_param.get("start") or 0) + 1
        res, total = query_address_traces(
            get_param,
            last=get_param.get("last") == "true",
            page_size=page_size,
            page_num=page_num,
        )
        return JsonResponse({"code": 200, "results": res, "count": total}, safe=False)

    def post(self, request):
        post_param = request.data
        result, _ = query_address_traces(
            post_param,
            last=bool(post_param.get("last")),
            page_size=int(post_param.get("page_size") or post_param.get("limit") or 1000),
            page_num=int(post_param.get("page") or 1),
        )

        columns = [
            {"label": "设备名称", "key": "device_name"},
            {"label": "机房", "key": "idc_name"},
            {"label": "设备序列号", "key": "device_serial_num"},
            {"label": "管理IP", "key": "manage_ip"},
            {"label": "接入口", "key": "interface_name"},
            {"label": "接入位置", "key": "node_location"},
            {"label": "目标IP", "key": "ip_address"},
            {"label": "MAC地址", "key": "mac_address"},
            {"label": "定位状态", "key": "trace_status"},
            {"label": "定位方式", "key": "trace_method"},
            {"label": "观测时间", "key": "observed_at"},
        ]

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Address location"
        for col_num, column in enumerate(columns, 1):
            worksheet[f"{get_column_letter(col_num)}1"] = column["label"]
        for row_num, device in enumerate(result, 2):
            for col_num, column in enumerate(columns, 1):
                worksheet[f"{get_column_letter(col_num)}{row_num}"] = device.get(column["key"], "")

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        filename = quote("地址寻觅.xlsx")
        response = HttpResponse(output, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f"attachment; filename*=UTF-8''{filename}"
        return response


class WorkflowSecurityView(APIView):
    def get(self, request):
        get_param = request.GET.dict()
        if get_param.get("get_firewall_list"):
            res = get_firewall_list()
            return JsonResponse({"code": 200, "count": 0, "message": "success", "results": res}, safe=False)
        return JsonResponse({"code": 400, "count": 0, "message": "没有匹配的数据", "results": []}, safe=False)

    def post(self, request):
        return JsonResponse({"code": 405, "message": "POST 未实现", "results": []}, safe=False)


class WorkflowDiagnoseView(APIView):
    def get(self, request):
        get_param = request.GET.dict()
        if "server_ip_address" not in get_param:
            return JsonResponse(data={})
        diagnose = WorkflowDiagnoseService(get_param["server_ip_address"])
        address_traces = diagnose.get_address_traces()
        if not address_traces:
            return JsonResponse(data={})
        manage_ip = address_traces[0]["manage_ip"]
        name = address_traces[0]["device_name"]
        cmdb_res = diagnose.get_cmdb(manage_ip)
        log_res = diagnose.get_log(manage_ip, name)
        lldp_res = diagnose.get_lldp(manage_ip)
        return JsonResponse(
            {
                "address_traces": address_traces,
                "cmdb": cmdb_res,
                "log": log_res,
                "lldp": lldp_res,
            }
        )


class WorkflowRuleToolView(APIView):
    permission_classes = ()
    authentication_classes = ()

    def get(self, request):
        get_param = request.GET.dict()
        if "get_cmdb_field" in get_param.keys():
            model = apps.get_model(app_label="asset", model_name="NetworkDevice")
            fields = model._meta.get_fields()
            field_res = []
            for field in fields:
                if isinstance(field, ForeignKey):
                    related_model = field.related_model
                    for related_field in related_model._meta.get_fields():
                        if isinstance(related_field, CharField):
                            if hasattr(related_field, "verbose_name"):
                                field_res.append(
                                    {
                                        "label": f"{field.verbose_name}-{related_field.verbose_name}",
                                        "value": f"{field.name}__{related_field.name}",
                                    }
                                )
                if isinstance(field, CharField) or isinstance(field, GenericIPAddressField):
                    field_res.append(
                        {
                            "label": getattr(field, "verbose_name", field.name),
                            "value": field.name,
                        }
                    )
            return JsonResponse({"code": 200, "msg": "success", "data": field_res})
        if "get_pulgin_list" in get_param.keys():
            return JsonResponse({"code": 200, "msg": "success", "data": auto_driver_map})
        return JsonResponse({"code": 400, "msg": "未匹配动作"})

    def post(self, request):
        return JsonResponse({"code": 400, "msg": "未匹配动作"})
